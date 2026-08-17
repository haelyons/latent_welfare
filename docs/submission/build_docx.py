#!/usr/bin/env python3
"""Populate the sprint submission template from the markdown draft.

    python build_docx.py <template.docx> <draft.md> <out.docx>

Keeps the template's title table (and its footnote), drops the guidance box and
every italicised instruction, and rebuilds the body from the markdown. Markdown
footnotes become real Word footnotes.
"""
import re
import sys
import zipfile
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

TEMPLATE, DRAFT, OUT = sys.argv[1], sys.argv[2], sys.argv[3]

AUTHOR = "Helios Lyons"
AFFIL = "Independent"

# --------------------------------------------------------------- md parsing
raw = open(DRAFT).read()

# split off the footnote definitions
notes = {}
body_md, _, notes_md = raw.partition("\n## Notes\n")
for m in re.finditer(r"^\[\^(\d+)\]:\s*(.*?)(?=\n\[\^\d+\]:|\Z)", notes_md, re.S | re.M):
    notes[int(m.group(1))] = " ".join(m.group(2).split())

# abstract + title
title = re.search(r"^# (.+)$", body_md, re.M).group(1).strip()
abstract = re.search(r"\*\*Abstract\*\*\s*\n(.*?)\n\s*---", body_md, re.S).group(1)
abstract = " ".join(abstract.split()).replace("this on on ", "this on ")

# everything from "## 1. Introduction" onward is the body
body_md = body_md[body_md.index("\n## 1. Introduction"):]

# ------------------------------------------------------------ block reader
def blocks(md):
    lines = md.split("\n")
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        if line.startswith("```"):
            j = i + 1
            buf = []
            while j < n and not lines[j].startswith("```"):
                buf.append(lines[j]); j += 1
            yield ("code", "\n".join(buf)); i = j + 1; continue
        if line.startswith("|"):
            buf = []
            while i < n and lines[i].startswith("|"):
                buf.append(lines[i]); i += 1
            yield ("table", buf); continue
        if line.startswith("### "):
            yield ("h3", line[4:].strip()); i += 1; continue
        if line.startswith("## "):
            yield ("h2", line[3:].strip()); i += 1; continue
        if line.startswith("> "):
            buf = []
            while i < n and lines[i].startswith(">"):
                buf.append(lines[i].lstrip("> ").rstrip()); i += 1
            yield ("quote", " ".join(x for x in buf if x)); continue
        if re.match(r"^[-*] ", line):
            buf = []
            while i < n and (re.match(r"^[-*] ", lines[i]) or (lines[i].startswith("  ") and lines[i].strip())):
                if re.match(r"^[-*] ", lines[i]):
                    buf.append(lines[i][2:].strip())
                else:
                    buf[-1] += " " + lines[i].strip()
                i += 1
            yield ("bullets", buf); continue
        if re.match(r"^\d+\. ", line):
            buf = []
            while i < n and (re.match(r"^\d+\. ", lines[i]) or (lines[i].startswith("   ") and lines[i].strip())):
                if re.match(r"^\d+\. ", lines[i]):
                    buf.append(re.sub(r"^\d+\.\s*", "", lines[i]).strip())
                else:
                    buf[-1] += " " + lines[i].strip()
                i += 1
            yield ("numbers", buf); continue
        if line.startswith("    ") and line.strip():
            buf = []
            while i < n and (lines[i].startswith("    ") or not lines[i].strip()):
                if not lines[i].strip():
                    if i + 1 < n and lines[i + 1].startswith("    "):
                        buf.append(""); i += 1; continue
                    break
                buf.append(lines[i][4:]); i += 1
            yield ("code", "\n".join(buf)); continue
        if not line.strip() or re.fullmatch(r"-{3,}", line.strip()):
            i += 1; continue
        buf = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{2,3} |\||```|> |[-*] |\d+\. )", lines[i]) \
                and not lines[i].startswith("    "):
            buf.append(lines[i].strip()); i += 1
        yield ("para", " ".join(buf))


# ------------------------------------------------------------ doc scaffold
doc = Document(TEMPLATE)
body = doc.element.body
children = list(body.iterchildren())
sectPr = children[-1]
title_tbl = children[1]

for ch in children:
    if ch is not title_tbl and ch is not sectPr:
        body.remove(ch)

# rebuild the title table cell
from docx.table import Table
t = Table(title_tbl, doc)
head, meta = t.rows[0].cells[0], t.rows[1].cells[0]


def wipe(cell):
    """Strip everything from a cell (paragraphs and nested tables), leave one empty p."""
    for ch in list(cell._element.iterchildren()):
        if ch.tag != qn("w:tcPr"):
            cell._element.remove(ch)
    return cell.add_paragraph()


def styled(par, text, *, bold=False, italic=False, size=None, mono=False):
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    r.font.name = "Courier New" if mono else "Old Standard TT"
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Courier New" if mono else "Old Standard TT")
    return r


p = wipe(head)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
styled(p, title, bold=True, size=20)

# keep the template's own title footnote (id 0: "Research conducted at ...")
_r = OxmlElement("w:r")
_rpr = OxmlElement("w:rPr")
_va = OxmlElement("w:vertAlign"); _va.set(qn("w:val"), "superscript")
_rpr.append(_va); _r.append(_rpr)
_ref = OxmlElement("w:footnoteReference")
_ref.set(qn("w:customMarkFollows"), "0"); _ref.set(qn("w:id"), "0")
_r.append(_ref)
p._element.append(_r)

p = wipe(meta)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
styled(p, AUTHOR, bold=True, size=11)
q = meta.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
styled(q, AFFIL, size=10)
q = meta.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
styled(q, "With", bold=True, size=10)
q = meta.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
styled(q, "Apart Research", size=10)
meta.add_paragraph()
q = meta.add_paragraph()
styled(q, "Abstract", bold=True, size=11)
abstract_par = meta.add_paragraph()
abstract_par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def new_par(style=None):
    par = doc.add_paragraph()
    if style:
        par.style = doc.styles[style]
    sectPr.addprevious(par._element)
    return par


# ---------------------------------------------------------- inline runs
FOOT = []  # (marker_run_element, note_number)
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|_[^_]+?_|`[^`]+?`|\[\^\d+\])")


def add_inline(par, text):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("[^"):
            num = int(tok[2:-1])
            r = par.add_run("")
            r.font.superscript = True
            FOOT.append((r._element, num))
        elif tok.startswith("**") and tok.endswith("**"):
            styled(par, tok[2:-2], bold=True, size=10)
        elif tok.startswith("`") and tok.endswith("`"):
            styled(par, tok[1:-1], mono=True, size=9)
        elif (tok.startswith("*") and tok.endswith("*")) or (tok.startswith("_") and tok.endswith("_")):
            styled(par, tok[1:-1], italic=True, size=10)
        else:
            styled(par, tok, size=10)


add_inline(abstract_par, abstract)


def set_borders(tbl):
    tblPr = tbl._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "999999")
        borders.append(e)
    tblPr.append(borders)


# ------------------------------------------------------------- emit body
for kind, payload in blocks(body_md):
    if kind == "h2":
        par = new_par("Heading 2")
        add_inline(par, payload)
    elif kind == "h3":
        par = new_par("Heading 3")
        add_inline(par, payload)
    elif kind == "para":
        par = new_par()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(par, payload)
    elif kind == "quote":
        par = new_par()
        par.paragraph_format.left_indent = Inches(0.4)
        par.paragraph_format.right_indent = Inches(0.4)
        styled(par, payload, italic=True, size=10)
    elif kind == "bullets":
        for b in payload:
            par = new_par()
            par.paragraph_format.left_indent = Inches(0.3)
            par.paragraph_format.first_line_indent = Inches(-0.18)
            styled(par, "•\t", size=10)
            add_inline(par, b)
    elif kind == "numbers":
        for k, b in enumerate(payload, 1):
            par = new_par()
            par.paragraph_format.left_indent = Inches(0.3)
            par.paragraph_format.first_line_indent = Inches(-0.22)
            styled(par, f"{k}.\t", size=10)
            add_inline(par, b)
    elif kind == "code":
        par = new_par()
        par.paragraph_format.left_indent = Inches(0.2)
        par.paragraph_format.space_after = Pt(6)
        for k, ln in enumerate(payload.split("\n")):
            if k:
                styled(par, "\n", mono=True, size=8)
            styled(par, ln, mono=True, size=8)
    elif kind == "table":
        rows = [r for r in payload if not re.match(r"^\|[\s|:-]+\|$", r)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        ncol = max(len(r) for r in cells)
        tbl = doc.add_table(rows=0, cols=ncol)
        tbl.style = doc.styles["TableNormal"]
        set_borders(tbl)
        for ri, row in enumerate(cells):
            wrow = tbl.add_row()
            for ci in range(ncol):
                cp = wrow.cells[ci].paragraphs[0]
                txt = row[ci] if ci < len(row) else ""
                if ri == 0:
                    styled(cp, re.sub(r"[*`]", "", txt), bold=True, size=8)
                else:
                    add_inline(cp, txt)
                    for r in cp.runs:
                        r.font.size = Pt(8)
        sectPr.addprevious(tbl._element)
        new_par()

doc.save(OUT)

# ------------------------------------------------- real footnotes via XML
NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{NS}}}"

zin = zipfile.ZipFile(OUT)
docx_xml = zin.read("word/document.xml").decode("utf-8")
fn_xml = zin.read("word/footnotes.xml").decode("utf-8")
others = {n: zin.read(n) for n in zin.namelist()
          if n not in ("word/document.xml", "word/footnotes.xml")}
zin.close()

# markers: our empty superscript runs, in document order
marker_re = re.compile(
    r'<w:r>(?:(?!</w:r>).)*?<w:vertAlign w:val="superscript"/>(?:(?!</w:r>).)*?</w:r>', re.S)
targets = [m for m in marker_re.finditer(docx_xml)]
# the template's own title footnote run also matches; skip runs already holding a reference
targets = [m for m in targets if "footnoteReference" not in m.group(0)]
assert len(targets) == len(FOOT), f"{len(targets)} markers vs {len(FOOT)} footnotes"

# Word numbers footnotes by reference order and expects one body per reference,
# so ids are assigned sequentially here and a repeated citation gets its own copy.
seq_bodies = []          # (new_id, text) in document order
out, last = [], 0
for k, (m, (_, num)) in enumerate(zip(targets, FOOT), start=1):
    seq_bodies.append((k, notes[num]))
    out.append(docx_xml[last:m.start()])
    out.append(
        f'<w:r><w:rPr><w:rFonts w:ascii="Old Standard TT" w:hAnsi="Old Standard TT" '
        f'w:cs="Old Standard TT" w:eastAsia="Old Standard TT"/><w:sz w:val="20"/>'
        f'<w:szCs w:val="20"/><w:vertAlign w:val="superscript"/></w:rPr>'
        f'<w:footnoteReference w:customMarkFollows="0" w:id="{k}"/></w:r>')
    last = m.end()
out.append(docx_xml[last:])
docx_xml = "".join(out)


def esc(s):
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


fn_new = []
for num, body_text in seq_bodies:
    txt = esc(body_text)
    fn_new.append(
        f'<w:footnote w:id="{num}"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
        f'w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Old Standard TT" '
        f'w:hAnsi="Old Standard TT" w:cs="Old Standard TT" w:eastAsia="Old Standard TT"/>'
        f'<w:sz w:val="16"/><w:szCs w:val="16"/></w:rPr></w:pPr>'
        f'<w:r><w:rPr><w:rFonts w:ascii="Old Standard TT" w:hAnsi="Old Standard TT" '
        f'w:cs="Old Standard TT" w:eastAsia="Old Standard TT"/><w:sz w:val="16"/>'
        f'<w:szCs w:val="16"/><w:vertAlign w:val="superscript"/></w:rPr>'
        f'<w:footnoteRef/></w:r>'
        f'<w:r><w:rPr><w:rFonts w:ascii="Old Standard TT" w:hAnsi="Old Standard TT" '
        f'w:cs="Old Standard TT" w:eastAsia="Old Standard TT"/><w:sz w:val="16"/>'
        f'<w:szCs w:val="16"/></w:rPr><w:t xml:space="preserve"> {txt}</w:t></w:r>'
        f'</w:p></w:footnote>')
fn_xml = fn_xml.replace("</w:footnotes>", "".join(fn_new) + "</w:footnotes>")

with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zo:
    for n, data in others.items():
        zo.writestr(n, data)
    zo.writestr("word/document.xml", docx_xml)
    zo.writestr("word/footnotes.xml", fn_xml)

print(f"wrote {OUT}: {len(FOOT)} footnote refs -> {len(seq_bodies)} bodies ({len(notes)} distinct notes)")
